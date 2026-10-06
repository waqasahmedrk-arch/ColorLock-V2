import csv
import io
import json

import pytest

from .test_auth import PASSWORD, sent  # noqa: F401  (fixture)
from .test_history import _score, _sign_in

pytestmark = pytest.mark.anyio


def _use(client, token):
    # Switch to another device's session cookie.
    client.cookies.clear()
    client.cookies.set("cl_session", token)

EMAIL = "a@example.com"


async def test_settings_endpoints_need_sign_in(client):
    assert (await client.get("/auth/sessions")).status_code == 401
    assert (await client.delete("/auth/sessions")).status_code == 401
    assert (await client.patch("/auth/me/preferences", json={"language": "zh"})).status_code == 401
    assert (await client.post("/auth/me/delete", json={"password": PASSWORD})).status_code == 401
    assert (await client.get("/history/export")).status_code == 401


async def test_sessions_list_and_sign_out(client, sent):  # noqa: F811
    await _sign_in(client, sent, EMAIL)
    other = client.cookies.get("cl_session")
    r = await client.post("/auth/login", json={"email": EMAIL, "password": PASSWORD},
                          headers={"User-Agent": "Mozilla/5.0 (Windows NT 10.0) Chrome/130.0"})
    assert r.status_code == 200
    mine = client.cookies.get("cl_session")

    sessions = (await client.get("/auth/sessions")).json()
    assert len(sessions) == 2
    assert sessions[0]["current"] and not sessions[1]["current"]
    assert "Chrome" in sessions[0]["user_agent"] and sessions[0]["ip"]
    assert sessions[0]["created_at"].endswith(("Z", "+00:00"))

    # Sign out the other device by id; this one stays signed in.
    assert (await client.delete(f"/auth/sessions/{sessions[1]['id']}")).status_code == 204
    assert (await client.delete("/auth/sessions/nope")).status_code == 404
    assert len((await client.get("/auth/sessions")).json()) == 1
    _use(client, other)
    assert (await client.get("/auth/me")).status_code == 401
    _use(client, mine)

    # "Sign out everywhere else" keeps only this session.
    r = await client.post("/auth/login", json={"email": EMAIL, "password": PASSWORD})
    newest = r.cookies.get("cl_session")
    _use(client, newest)
    assert (await client.delete("/auth/sessions")).status_code == 204
    assert [s["current"] for s in (await client.get("/auth/sessions")).json()] == [True]
    _use(client, mine)
    assert (await client.get("/auth/me")).status_code == 401
    _use(client, newest)
    assert (await client.get("/auth/me")).status_code == 200


async def test_language_is_saved_and_restored_at_login(client, sent):  # noqa: F811
    await _sign_in(client, sent, EMAIL)
    assert (await client.get("/auth/me")).json()["language"] is None
    r = await client.patch("/auth/me/preferences", json={"language": "zh"})
    assert r.status_code == 200 and r.json()["language"] == "zh"
    assert r.cookies.get("lang") == "zh"
    assert (await client.patch("/auth/me/preferences", json={"language": "fr"})).status_code == 422

    client.cookies.clear()
    r = await client.post("/auth/login", json={"email": EMAIL, "password": PASSWORD})
    assert r.status_code == 200 and r.cookies.get("lang") == "zh"


async def test_export_history_csv_and_json(client, sent):  # noqa: F811
    await _sign_in(client, sent, EMAIL)
    await _score(client, "=cmd.png", score_name="=SUM(A1)")
    await _score(client, "second.png", score_name="Second")

    r = await client.get("/history/export")
    assert r.status_code == 200 and r.headers["content-type"].startswith("text/csv")
    assert 'filename="colorlock-history-' in r.headers["content-disposition"]
    rows = list(csv.reader(io.StringIO(r.text)))
    assert rows[0][:4] == ["created_at", "name", "filename", "target_id"] and len(rows) == 3
    assert rows[1][1] == "Second" and rows[2][1] == "'=SUM(A1)"
    assert rows[1][2] == "second.png" and rows[2][2] == "'=cmd.png"
    assert rows[1][3] == "teal" and rows[1][8] == "yes"

    r = await client.get("/history/export?format=json")
    data = json.loads(r.text)
    assert r.headers["content-type"].startswith("application/json") and len(data) == 2
    assert data[0]["result"]["target"]["id"] == "teal"
    assert (await client.get("/history/export?format=xml")).status_code == 422


async def test_delete_account(client, sent):  # noqa: F811
    await _sign_in(client, sent, EMAIL)
    await _score(client)
    r = await client.post("/auth/me/delete", json={"password": "wrong-pass-1"})
    assert r.status_code == 400 and r.json()["title"] == "Incorrect password"
    assert (await client.get("/auth/me")).status_code == 200

    assert (await client.post("/auth/me/delete", json={"password": PASSWORD})).status_code == 204
    assert (await client.get("/auth/me")).status_code == 401
    r = await client.post("/auth/login", json={"email": EMAIL, "password": PASSWORD})
    assert r.status_code == 401
    # The address is free to sign up again, with an empty history.
    client.cookies.clear()
    await _sign_in(client, sent, EMAIL)
    assert (await client.get("/history")).json()["total"] == 0


async def test_export_history_xlsx(client, sent):  # noqa: F811
    from openpyxl import load_workbook

    await _sign_in(client, sent, EMAIL)
    await _score(client, "=cmd.png", score_name="=HYPERLINK(1)")
    await _score(client, "second.png", score_name="Second")

    r = await client.get("/history/export?format=xlsx")
    assert r.status_code == 200
    assert r.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    assert r.headers["content-disposition"].endswith('.xlsx"')
    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["Scores", "About"]
    ws = wb["Scores"]
    head = [c.value for c in ws[1]]
    assert head[:3] == ["scored_at_utc", "name", "filename"] and "delta_e00" in head
    assert ws.max_row == 3
    row = {h: c for h, c in zip(head, ws[2], strict=True)}  # newest first
    assert row["name"].value == "Second" and row["target_id"].value == "teal"
    assert isinstance(row["delta_e00"].value, float) and row["qc_pass"].value is True
    assert isinstance(row["sample_L"].value, float) and row["package_version"].value
    # User text that looks like a formula stays text.
    first = {h: c for h, c in zip(head, ws[3], strict=True)}
    assert first["name"].value == "=HYPERLINK(1)" and first["name"].data_type == "s"
    assert first["filename"].value == "=cmd.png" and first["filename"].data_type == "s"
    about = [c.value for c in wb["About"]["A"]]
    assert "Scores" in about and "delta_e00" in about
