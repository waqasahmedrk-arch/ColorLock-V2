"""Score history as an Excel workbook, for users collecting data for their own research.

Sheet "Scores" has one row per saved score with every measured value (typed numbers, so it
can be filtered, charted and analysed directly); sheet "About" says what each column means and
how the values were measured.
"""

from __future__ import annotations

import io
from collections.abc import Callable, Sequence
from datetime import UTC, datetime
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from ..db.auth import ScoreRecord, User

MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

_HEAD_FONT = Font(bold=True, color="FFFFFF")
_HEAD_FILL = PatternFill("solid", fgColor="1B1B19")
_NUM = "0.0000"


def _lab(v: Any, i: int) -> float | None:
    try:
        return float(v[i])
    except (TypeError, IndexError, ValueError):
        return None


# (header, number format or None, unit, description, value getter)
Column = tuple[str, str | None, str, str, Callable[[ScoreRecord, dict[str, Any]], Any]]

COLUMNS: list[Column] = [
    ("scored_at_utc", "yyyy-mm-dd hh:mm:ss", "UTC", "When the image was scored.",
     lambda r, _: r.created_at),
    ("name", None, "", "Your name for the score (empty for scores saved before names existed).",
     lambda r, _: r.name),
    ("filename", None, "", "Original file name of the uploaded image (the image itself is not kept).",
     lambda r, _: r.filename),
    ("target_id", None, "", "Study colour id, or empty for a custom hex target.",
     lambda r, _: r.target_id),
    ("target_name", None, "", "Target colour name.", lambda r, _: r.target_name),
    ("target_hex", None, "sRGB", "Target colour.", lambda r, _: r.target_hex),
    ("target_L", _NUM, "CIELAB", "Target L*.", lambda _, x: _lab(x.get("target", {}).get("lab"), 0)),
    ("target_a", _NUM, "CIELAB", "Target a*.", lambda _, x: _lab(x.get("target", {}).get("lab"), 1)),
    ("target_b", _NUM, "CIELAB", "Target b*.", lambda _, x: _lab(x.get("target", {}).get("lab"), 2)),
    ("sample_hex", None, "sRGB", "Measured colour (dominant colour of the central crop).",
     lambda r, _: r.sample_hex),
    ("sample_L", _NUM, "CIELAB", "Measured L*.", lambda _, x: _lab(x.get("sample_lab"), 0)),
    ("sample_a", _NUM, "CIELAB", "Measured a*.", lambda _, x: _lab(x.get("sample_lab"), 1)),
    ("sample_b", _NUM, "CIELAB", "Measured b*.", lambda _, x: _lab(x.get("sample_lab"), 2)),
    ("delta_e00", _NUM, "ΔE00", "CIEDE2000 difference between measured and target colour; lower is closer.",
     lambda r, _: r.delta_e00),
    ("flat_p95_de", _NUM, "ΔE00", "Flatness: 95th percentile ΔE00 of crop pixels from their median.",
     lambda _, x: x.get("flat_p95_de")),
    ("qc_threshold", _NUM, "ΔE00", "Flatness QC threshold in force when this score was made.",
     lambda _, x: x.get("qc", {}).get("threshold")),
    ("qc_pass", None, "TRUE/FALSE", "Whether flat_p95_de ≤ qc_threshold (a flat colour field).",
     lambda r, _: bool(r.qc_pass)),
    ("chroma_sample", _NUM, "C*", "Chroma of the measured colour.",
     lambda _, x: x.get("chroma", {}).get("sample")),
    ("chroma_reference", _NUM, "C*", "Chroma of the target colour.",
     lambda _, x: x.get("chroma", {}).get("reference")),
    ("chroma_delta", _NUM, "C*", "chroma_sample − chroma_reference; negative = less saturated.",
     lambda _, x: x.get("chroma", {}).get("delta")),
    ("kept_pct", "0.00", "%", "Legacy foreground-share metric (not used for QC).",
     lambda _, x: x.get("kept_pct")),
    ("warnings", None, "", "Input warnings, separated by '; ' (e.g. jpeg_input, alpha_flattened).",
     lambda _, x: "; ".join(x.get("warnings") or [])),
    ("package_version", None, "", "Version of the colourlock measurement package that scored it.",
     lambda _, x: x.get("package_version")),
    ("score_id", None, "", "Score id (for reference / support).", lambda r, _: r.score_id),
]
_WIDTHS = {"scored_at_utc": 20, "name": 28, "filename": 26, "target_name": 18, "warnings": 28,
           "score_id": 38, "package_version": 16}


def _text(cell: Any) -> None:
    # openpyxl stores text beginning with "=" as a formula; keep user text as text.
    if isinstance(cell.value, str):
        cell.data_type = "s"


def build(rows: Sequence[ScoreRecord], user: User) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Scores"
    for c, (head, *_rest) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=c, value=head)
        cell.font, cell.fill = _HEAD_FONT, _HEAD_FILL
        cell.alignment = Alignment(vertical="center")
    for r_i, rec in enumerate(rows, start=2):
        result = rec.result or {}
        for c, (_head, fmt, _unit, _desc, get) in enumerate(COLUMNS, start=1):
            cell = ws.cell(row=r_i, column=c, value=get(rec, result))
            _text(cell)
            if fmt:
                cell.number_format = fmt
    for c, (head, *_rest) in enumerate(COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(c)].width = _WIDTHS.get(head, 14)
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{max(len(rows) + 1, 1)}"

    about = wb.create_sheet("About")
    versions = sorted({str((r.result or {}).get("package_version")) for r in rows} - {"None"})
    info = [
        ("ColorLock score history export", None),
        ("Exported at (UTC)", datetime.now(UTC).replace(tzinfo=None)),
        ("Account", f"{user.name} <{user.email}>"),
        ("Scores", len(rows)),
        ("Measured with", ", ".join(f"colourlock {v}" for v in versions) or "—"),
        ("Method", ("Each image's dominant colour in the central crop is compared with the target "
                    "using CIEDE2000 (ΔE00). The flatness QC flags images that are not a flat colour "
                    "field; the study only counts QC-passed images for accuracy.")),
        ("Note", ("Each row is a single-image measurement. Consistency and the study's hypothesis "
                  "tests apply only to the frozen study dataset and are not part of this export.")),
        ("Note", "Timestamps are UTC. Hex values are sRGB; L*, a*, b* are CIELAB (D65)."),
    ]
    for r_i, (key, value) in enumerate(info, start=1):
        about.cell(row=r_i, column=1, value=key).font = Font(bold=True, size=13 if r_i == 1 else 11)
        cell = about.cell(row=r_i, column=2, value=value)
        _text(cell)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        if isinstance(value, datetime):
            cell.number_format = "yyyy-mm-dd hh:mm:ss"
    start = len(info) + 2
    for c, head in enumerate(("Column", "Unit", "Description"), start=1):
        cell = about.cell(row=start, column=c, value=head)
        cell.font, cell.fill = _HEAD_FONT, _HEAD_FILL
    for i, (head, _fmt, unit, desc, _get) in enumerate(COLUMNS, start=start + 1):
        about.cell(row=i, column=1, value=head).font = Font(name="Consolas")
        about.cell(row=i, column=2, value=unit)
        about.cell(row=i, column=3, value=desc).alignment = Alignment(wrap_text=True, vertical="top")
    about.column_dimensions["A"].width = 20
    about.column_dimensions["B"].width = 60
    about.column_dimensions["C"].width = 70

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
